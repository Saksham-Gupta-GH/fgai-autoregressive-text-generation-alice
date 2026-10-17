import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def add_docx_hyperlink(paragraph, url, text, color="0066CC", underline=True):
    """
    Inserts a genuine clickable hyperlink into a python-docx paragraph.
    """
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def build_submission_docx(
    student_name="Saksham Gupta",
    reg_num="230911186",
    repo_url="https://github.com/Saksham-Gupta-GH/fgai-autoregressive-text-generation-alice",
    output_filename="SakshamGupta_230911186_FGAI_Assignment2.docx"
):
    doc = Document()
    
    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("FOUNDATIONS OF GENERATIVE AI (ICT4403)")
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)
    
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub1 = sub.add_run("Assignment 2: Autoregressive Text Generation System\n")
    run_sub1.font.size = Pt(12)
    run_sub1.font.bold = True
    run_sub2 = sub.add_run("Submission Document & GitHub Repository Verification Sheet")
    run_sub2.font.size = Pt(10)
    run_sub2.font.italic = True
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Section 1: Required Details
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Required Submission Details (Step 3)")
    r_h1.font.size = Pt(11.5)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0, 0, 0)
    
    # Metadata Table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    meta_rows = [
        ("Student Name", student_name, None),
        ("Registration Number", reg_num, None),
        ("Assignment Title", "Foundations of Generative AI (ICT4403) — Assignment 2", None),
        ("Title of the Developed Autoregressive Text-Generation System", "Comparative Word-Level Autoregressive Text Generation on Lewis Carroll's Alice's Adventures in Wonderland using Simple RNN, Bidirectional RNN, LSTM, and GRU", None),
        ("Name and Link of the Dataset Used", "Alice's Adventures in Wonderland (Lewis Carroll, Project Gutenberg eBook #11)\nDataset Link: ", "https://www.gutenberg.org/files/11/11-0.txt"),
        ("GitHub Repository Link", "Repository Link: ", repo_url)
    ]
    
    for i, (k, text_prefix, link_url) in enumerate(meta_rows):
        row = table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.5)
        c1.width = Inches(4.4)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        
        p1 = c1.paragraphs[0]
        if text_prefix:
            r1 = p1.add_run(text_prefix)
            r1.font.size = Pt(9.5)
        if link_url:
            add_docx_hyperlink(p1, link_url, link_url)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section 2: Repository Structure & Deliverables
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. GitHub Repository Structure & Verified Deliverables (Step 2)")
    r_h2.font.size = Pt(11.5)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0, 0, 0)
    
    p_repo_intro = doc.add_paragraph(
        "All implementation code, pre-trained model checkpoints, interactive notebooks, plots, and experimental data "
        "are hosted in the public GitHub repository above in strict compliance with the assignment guidelines:"
    )
    p_repo_intro.paragraph_format.space_after = Pt(4)
    
    items = [
        ("Complete Source Implementation (src/): ", "Modular PyTorch implementation containing preprocessing.py (chronological 80/10/10 split, train-only vocabulary fitting with min_freq=2), models.py (Simple RNN, BiRNN with [h_n[-2], h_n[-1]] state concatenation, LSTM, and GRU), train.py (training loop with validation checkpointing), evaluate.py (loss, perplexity, Top-1/Top-5 accuracy), and generate.py (temperature sampling, Distinct-1/2, trigram repetition metrics)."),
        ("Interactive Jupyter Notebook (notebooks/text_generation_assignment2.ipynb): ", "Complete, fully executed notebook with all 11 cells containing data preparation, model training progress, evaluation plots, and generated text outputs."),
        ("Automated End-to-End Pipeline (run_experiments.py): ", "Single-command reproduction script executing the full experimental workflow: preprocessing, training, evaluation, plotting, and generation."),
        ("Empirical Results & Visualizations (results/): ", "Tabular CSV reports, generated sample logs across temperatures (T ∈ {0.2, 0.5, 1.0, 1.5}), and high-resolution comparison plots (loss_curves.png, perplexity_comparison.png, accuracy_comparison.png, training_time_comparison.png)."),
        ("Comprehensive Technical Report (README.md): ", "Complete report covering problem formulation, dataset characteristics, architecture details, experimental protocol, full evaluation tables, temperature analysis, and model justification.")
    ]
    
    for prefix, body in items:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(3)
        r_pre = p_item.add_run(prefix)
        r_pre.font.bold = True
        r_pre.font.size = Pt(9)
        r_b = p_item.add_run(body)
        r_b.font.size = Pt(9)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # Section 3: Summary of Key Findings
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Summary of Key Experimental Findings")
    r_h3.font.size = Pt(11.5)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0, 0, 0)
    
    findings = [
        ("Best Performing Model: ", "LSTM achieved the strongest next-word prediction performance (Best Val Loss: 3.9903, Val Perplexity: 54.07, Held-Out Test Top-1 Accuracy: 23.29%, Test Top-5 Accuracy: 47.23%), with GRU performing closely behind (Val Loss: 4.0353, Test Top-1: 22.53%) with 11% fewer parameters."),
        ("Sampling Temperature Optimization: ", "Temperature T = 0.5 demonstrated the optimal empirical trade-off between lexical diversity (Distinct-1: 0.60–0.73) and structural coherence, eliminating repetition loops while maintaining fluent narrative dialogue tags.")
    ]
    
    for prefix, body in findings:
        p_f = doc.add_paragraph(style='List Bullet')
        p_f.paragraph_format.space_after = Pt(3)
        r_p = p_f.add_run(prefix)
        r_p.font.bold = True
        r_p.font.size = Pt(9)
        r_b = p_f.add_run(body)
        r_b.font.size = Pt(9)
        
    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_footer.add_run("Submitted in fulfillment of Steps 3 and 4 of the Assignment 2 submission procedure for ICT4403.")
    r_foot.font.italic = True
    r_foot.font.size = Pt(8.5)
    
    output_path = os.path.join(os.getcwd(), output_filename)
    doc.save(output_path)
    print(f"Generated Word document at: {output_path}")
    return output_path


def build_submission_pdf(
    student_name="Saksham Gupta",
    reg_num="230911186",
    repo_url="https://github.com/Saksham-Gupta-GH/fgai-autoregressive-text-generation-alice",
    output_filename="SakshamGupta_230911186_FGAI_Assignment2.pdf"
):
    pdf_path = os.path.join(os.getcwd(), output_filename)
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=24, bottomMargin=24)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=13, leading=16, textColor=colors.black, alignment=1, spaceAfter=2)
    sub_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=9, leading=12, fontName='Helvetica-Oblique', alignment=1, spaceAfter=6)
    h1_style = ParagraphStyle('SecH1', parent=styles['Heading2'], fontSize=10, leading=13, textColor=colors.black, spaceBefore=5, spaceAfter=3)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=7.5, leading=10, spaceAfter=4)
    table_text = ParagraphStyle('TableT', parent=styles['Normal'], fontSize=7.5, leading=9.5)
    table_head = ParagraphStyle('TableH', parent=styles['Normal'], fontSize=7.5, leading=9.5, fontName='Helvetica-Bold')

    story = []
    story.append(Paragraph("FOUNDATIONS OF GENERATIVE AI (ICT4403)", title_style))
    story.append(Paragraph("Assignment 2: Autoregressive Text Generation System — Comparative Study", sub_style))

    # Meta Table
    meta_table_data = [
        [Paragraph("<b>Student Name</b>", table_text), Paragraph(student_name, table_text)],
        [Paragraph("<b>Registration Number</b>", table_text), Paragraph(reg_num, table_text)],
        [Paragraph("<b>Assignment Title</b>", table_text), Paragraph("Assignment 2 - Autoregressive Text Generation System", table_text)],
        [Paragraph("<b>System Title</b>", table_text), Paragraph("Comparative Word-Level Autoregressive Text Generation on Lewis Carroll's Alice in Wonderland using Simple RNN, BiRNN, LSTM, GRU", table_text)],
        [Paragraph("<b>Dataset & Source</b>", table_text), Paragraph('Alice\'s Adventures in Wonderland (Lewis Carroll) | <a href="https://www.gutenberg.org/files/11/11-0.txt"><font color="#0066cc"><u>https://www.gutenberg.org/files/11/11-0.txt</u></font></a>', table_text)],
        [Paragraph("<b>GitHub Repository</b>", table_text), Paragraph(f'<a href="{repo_url}"><font color="#0066cc"><u>{repo_url}</u></font></a>', table_text)],
    ]
    meta_table = Table(meta_table_data, colWidths=[120, 420])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8f9fa')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cccccc')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ]))
    story.append(meta_table)

    # Section 1
    story.append(Paragraph("1. System Overview & Problem Statement", h1_style))
    story.append(Paragraph(
        "Four recurrent neural language models (Simple RNN, BiRNN with forward-backward final state concatenation [h_n[-2], h_n[-1]], LSTM, and GRU) were trained "
        "under strictly controlled conditions on Alice's Adventures in Wonderland. To eliminate data leakage, the corpus is partitioned chronologically (80% Train, 10% "
        "Val, 10% Test) and the vocabulary (1,300 tokens, min_freq=2) is fitted exclusively on training tokens. Context sequence length is 6 tokens. Held-out test set has an 11.17% &lt;unk&gt; token rate.",
        body_style
    ))

    # Section 2
    story.append(Paragraph("2. Quantitative Evaluation Summary (Best Epoch & Test Set)", h1_style))
    q_data = [
        [Paragraph(h, table_head) for h in ["Architecture", "Parameters", "Best Ep", "Val Loss", "Val PPL", "Val Top-1", "Test Loss", "Test Top-1", "Test Top-5"]],
        [Paragraph("<b>Simple RNN</b>", table_text), Paragraph("798,484", table_text), Paragraph("4", table_text), Paragraph("4.0741", table_text), Paragraph("58.80", table_text), Paragraph("19.73%", table_text), Paragraph("4.1584", table_text), Paragraph("20.80%", table_text), Paragraph("45.28%", table_text)],
        [Paragraph("<b>BiRNN</b>", table_text), Paragraph("1,262,868", table_text), Paragraph("2", table_text), Paragraph("4.2804", table_text), Paragraph("72.27", table_text), Paragraph("18.66%", table_text), Paragraph("4.3192", table_text), Paragraph("18.47%", table_text), Paragraph("42.26%", table_text)],
        [Paragraph("<b>LSTM</b>", table_text), Paragraph("1,193,236", table_text), Paragraph("5", table_text), Paragraph("<b>3.9903</b>", table_text), Paragraph("<b>54.07</b>", table_text), Paragraph("<b>21.87%</b>", table_text), Paragraph("<b>4.0576</b>", table_text), Paragraph("<b>23.29%</b>", table_text), Paragraph("<b>47.23%</b>", table_text)],
        [Paragraph("<b>GRU</b>", table_text), Paragraph("1,061,652", table_text), Paragraph("3", table_text), Paragraph("4.0353", table_text), Paragraph("56.56", table_text), Paragraph("20.55%", table_text), Paragraph("4.0900", table_text), Paragraph("22.53%", table_text), Paragraph("46.51%", table_text)],
    ]
    q_table = Table(q_data, colWidths=[75, 65, 42, 50, 50, 52, 50, 56, 56])
    q_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e9ecef')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cccccc')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ]))
    story.append(q_table)

    # Section 3
    story.append(Paragraph("3. Cross-Model Temperature Analysis (Averaged Over Models and Seed Prompts, n = 3 per cell)", h1_style))
    t_data = [
        [Paragraph(h, table_head) for h in ["T", "Distinct-1", "Distinct-2", "Trigram Repetition", "Longest Loop", "Empirical Quality Observation"]],
        [Paragraph("<b>0.2</b>", table_text), Paragraph("0.517", table_text), Paragraph("0.674", table_text), Paragraph("23.7%", table_text), Paragraph("0–6 words", table_text), Paragraph("Strict syntax; repetitive looping traps across models (GRU 52.5%)", table_text)],
        [Paragraph("<b>0.5</b>", table_text), Paragraph("0.678", table_text), Paragraph("0.900", table_text), Paragraph("3.0%", table_text), Paragraph("0 words", table_text), Paragraph("Optimal empirical trade-off: high diversity with fluent story dialogue; LSTM and GRU comparable", table_text)],
        [Paragraph("<b>1.0</b>", table_text), Paragraph("0.855", table_text), Paragraph("0.992", table_text), Paragraph("0.0%", table_text), Paragraph("0 words", table_text), Paragraph("High lexical richness; slight syntactic drift across longer clauses", table_text)],
        [Paragraph("<b>1.5</b>", table_text), Paragraph("0.916", table_text), Paragraph("1.000", table_text), Paragraph("0.0%", table_text), Paragraph("0 words", table_text), Paragraph("High entropy; syntactic collapse into ungrammatical token sequences", table_text)],
    ]
    t_table = Table(t_data, colWidths=[35, 52, 52, 72, 55, 274])
    t_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e9ecef')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cccccc')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ]))
    story.append(t_table)

    # Section 4
    story.append(Paragraph("4. Conclusion & Model Comparison", h1_style))
    story.append(Paragraph(
        "• <b>LSTM vs. GRU & UNK Nuance:</b> LSTM leads across next-word prediction metrics (Val Loss 3.9903, PPL 54.07, Test Top-1 23.29%), with GRU performing closely behind (Val Loss 4.0353, Test Top-1 22.53%) with 11% fewer parameters. Test targets include 11.17% &lt;unk&gt; tokens due to train-only vocabulary fitting.<br/>"
        "• <b>Qualitative Consistency & Seed Origin:</b> Seed 1 comes from the training text, Seed 2 from held-out test text, and Seed 3 from validation text. Consistency was assessed qualitatively: at T=0.5, LSTM and GRU more frequently generate Carroll-style dialogue tags ('the march hare said'), though all models drift across multi-clause sequences without sustaining a persistent scene.<br/>"
        "• <b>Single-Seed Limitation & BiRNN:</b> SEED=42 ensures strict comparability; the 0.045 loss gap between LSTM and GRU falls within single-run variance margins. BiRNN's forward-backward state concatenation feeds complete context, but its 1,262,868 parameters overfit by epoch 2 (val loss 4.2804).<br/>"
        "• <b>Temperature Sampling:</b> T = 0.5 provides the optimal trade-off between lexical diversity (Distinct-1: 0.60–0.73) and phrase stability, achieving minimal trigram repetition (LSTM 5.0%, GRU 0.0%, BiRNN 1.0%, Simple RNN 6.1%).",
        body_style
    ))

    doc.build(story)
    print(f"Generated PDF document at: {pdf_path}")
    return pdf_path


if __name__ == "__main__":
    docx_file = build_submission_docx()
    pdf_file = build_submission_pdf()
